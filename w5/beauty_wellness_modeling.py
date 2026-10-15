"""Reproducible customer-return modeling example for beauty and wellness businesses.

With no input file, the script generates a clearly labeled synthetic dataset.
To use real data, provide a CSV with the documented feature columns and the
``returned_within_60d`` binary target.
"""

from __future__ import annotations

import argparse
import json
import logging
from pathlib import Path
from typing import Any

import joblib
import matplotlib

matplotlib.use("Agg")

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt


LOGGER = logging.getLogger("beauty_wellness_modeling")
TARGET = "returned_within_60d"
NUMERIC_FEATURES = [
    "visits_last_90d",
    "avg_ticket_usd",
    "days_since_last_visit",
    "discount_share",
    "satisfaction_score",
]
CATEGORICAL_FEATURES = ["membership", "service_category"]
FEATURES = NUMERIC_FEATURES + CATEGORICAL_FEATURES
SYNTHETIC_ROWS = 3000
RANDOM_STATE = 42


def generate_synthetic_data(
    n_rows: int = SYNTHETIC_ROWS, random_state: int = RANDOM_STATE
) -> pd.DataFrame:
    """Generate illustrative, non-customer data with a known return signal."""
    if n_rows < 100:
        raise ValueError("Synthetic dataset size must be at least 100 rows.")

    rng = np.random.default_rng(random_state)
    visits = np.clip(rng.poisson(2.5, n_rows), 0, 12)
    ticket = np.clip(rng.lognormal(np.log(68), 0.42, n_rows), 20, 250)
    recency = np.clip(rng.gamma(2.0, 24.0, n_rows), 1, 180)
    membership = rng.choice(["member", "non_member"], n_rows, p=[0.38, 0.62])
    discount_share = np.clip(rng.beta(2.0, 5.0, n_rows), 0, 1)
    satisfaction = np.clip(
        np.rint(rng.normal(4.0, 0.8, n_rows)), 1, 5
    ).astype(int)
    service = rng.choice(
        ["hair", "skincare", "massage", "nails"],
        n_rows,
        p=[0.30, 0.28, 0.22, 0.20],
    )

    log_odds = (
        -0.55
        + 0.23 * visits
        - 0.017 * recency
        + 0.32 * satisfaction
        + 0.62 * (membership == "member")
        - 0.75 * discount_share
        + 0.0018 * (ticket - 68)
        + 0.25 * (service == "skincare")
        - 0.12 * (service == "nails")
    )
    probability = 1 / (1 + np.exp(-log_odds))
    returned = rng.binomial(1, probability)

    return pd.DataFrame(
        {
            "visits_last_90d": visits,
            "avg_ticket_usd": np.round(ticket, 2),
            "days_since_last_visit": np.round(recency).astype(int),
            "membership": membership,
            "discount_share": np.round(discount_share, 3),
            "satisfaction_score": satisfaction,
            "service_category": service,
            TARGET: returned,
        }
    )


def load_data(input_path: Path | None) -> tuple[pd.DataFrame, bool]:
    """Load a supplied CSV or create synthetic data, with clear schema errors."""
    if input_path is None:
        LOGGER.info("No CSV supplied; generating illustrative synthetic data.")
        return generate_synthetic_data(), True

    if not input_path.is_file():
        raise FileNotFoundError(f"Input CSV not found: {input_path}")
    frame = pd.read_csv(input_path)
    missing = [column for column in FEATURES + [TARGET] if column not in frame]
    if missing:
        raise ValueError(
            "Input CSV is missing required columns: " + ", ".join(missing)
        )
    frame = frame[FEATURES + [TARGET]].copy()
    if frame.isna().any().any():
        missing_counts = frame.isna().sum()
        details = ", ".join(
            f"{name}={int(count)}"
            for name, count in missing_counts.items()
            if count > 0
        )
        raise ValueError(
            "Input data contains missing values; resolve them before training "
            f"so missingness is handled deliberately. Counts: {details}"
        )
    return frame, False


def validate_data(frame: pd.DataFrame) -> None:
    """Fail early for unusable labels, small samples, or invalid feature values."""
    if len(frame) < 100:
        raise ValueError("At least 100 complete observations are required.")
    if frame[TARGET].nunique() != 2 or set(frame[TARGET].unique()) != {0, 1}:
        raise ValueError(f"{TARGET} must contain both binary labels 0 and 1.")
    for column in NUMERIC_FEATURES:
        if not pd.api.types.is_numeric_dtype(frame[column]):
            raise TypeError(f"Feature {column!r} must be numeric.")
        if not np.isfinite(frame[column].to_numpy(dtype=float)).all():
            raise ValueError(f"Feature {column!r} contains non-finite values.")


def build_models() -> dict[str, Any]:
    """Create comparable preprocessing/model pipelines."""
    try:
        from sklearn.compose import ColumnTransformer
        from sklearn.ensemble import RandomForestClassifier
        from sklearn.linear_model import LogisticRegression
        from sklearn.pipeline import Pipeline
        from sklearn.preprocessing import OneHotEncoder, StandardScaler
    except ImportError as error:
        raise RuntimeError(
            "scikit-learn could not be loaded. Install the dependencies from "
            "requirements.txt and check whether Windows security policy is "
            "blocking its compiled extensions."
        ) from error

    logistic_preprocessor = ColumnTransformer(
        [
            ("numeric", StandardScaler(), NUMERIC_FEATURES),
            (
                "categorical",
                OneHotEncoder(handle_unknown="ignore"),
                CATEGORICAL_FEATURES,
            ),
        ]
    )
    tree_preprocessor = ColumnTransformer(
        [
            ("numeric", "passthrough", NUMERIC_FEATURES),
            (
                "categorical",
                OneHotEncoder(handle_unknown="ignore"),
                CATEGORICAL_FEATURES,
            ),
        ]
    )
    return {
        "Logistic regression": Pipeline(
            [
                ("preprocess", logistic_preprocessor),
                (
                    "model",
                    LogisticRegression(
                        class_weight="balanced",
                        max_iter=1500,
                        random_state=RANDOM_STATE,
                    ),
                ),
            ]
        ),
        "Random forest": Pipeline(
            [
                ("preprocess", tree_preprocessor),
                (
                    "model",
                    RandomForestClassifier(
                        n_estimators=300,
                        min_samples_leaf=8,
                        max_features="sqrt",
                        class_weight="balanced_subsample",
                        random_state=RANDOM_STATE,
                        n_jobs=-1,
                    ),
                ),
            ]
        ),
    }


def choose_f1_threshold(
    labels: np.ndarray, probabilities: np.ndarray
) -> tuple[float, float]:
    """Choose an operating threshold from out-of-fold training predictions."""
    from sklearn.metrics import f1_score

    thresholds = np.arange(0.10, 0.81, 0.01)
    scores = [
        f1_score(labels, probabilities >= threshold, zero_division=0)
        for threshold in thresholds
    ]
    best_index = int(np.argmax(scores))
    return float(thresholds[best_index]), float(scores[best_index])


def evaluate(
    labels: np.ndarray, probabilities: np.ndarray, threshold: float
) -> dict[str, float | int]:
    from sklearn.metrics import (
        accuracy_score,
        average_precision_score,
        brier_score_loss,
        confusion_matrix,
        f1_score,
        precision_score,
        recall_score,
        roc_auc_score,
    )

    predicted = (probabilities >= threshold).astype(int)
    tn, fp, fn, tp = confusion_matrix(labels, predicted, labels=[0, 1]).ravel()
    return {
        "roc_auc": float(roc_auc_score(labels, probabilities)),
        "pr_auc_average_precision": float(
            average_precision_score(labels, probabilities)
        ),
        "precision": float(precision_score(labels, predicted, zero_division=0)),
        "recall": float(recall_score(labels, predicted, zero_division=0)),
        "f1": float(f1_score(labels, predicted, zero_division=0)),
        "accuracy": float(accuracy_score(labels, predicted)),
        "specificity": float(tn / (tn + fp)) if tn + fp else 0.0,
        "brier_score": float(brier_score_loss(labels, probabilities)),
        "true_negative": int(tn),
        "false_positive": int(fp),
        "false_negative": int(fn),
        "true_positive": int(tp),
    }


def make_roc_plot(
    y_test: pd.Series,
    probabilities: dict[str, np.ndarray],
    output_path: Path,
) -> None:
    from sklearn.metrics import roc_auc_score, roc_curve

    fig, ax = plt.subplots(figsize=(7.2, 5.0))
    for name, scores in probabilities.items():
        false_positive_rate, true_positive_rate, _ = roc_curve(y_test, scores)
        auc = roc_auc_score(y_test, scores)
        ax.plot(false_positive_rate, true_positive_rate, label=f"{name} (AUC={auc:.3f})")
    ax.plot([0, 1], [0, 1], linestyle="--", color="gray", label="No-skill baseline")
    ax.set(
        xlabel="False positive rate",
        ylabel="True positive rate",
        title="Held-out ROC curves",
        xlim=(0, 1),
        ylim=(0, 1.02),
    )
    ax.legend(loc="lower right")
    ax.grid(alpha=0.2)
    fig.tight_layout()
    fig.savefig(output_path, dpi=160)
    plt.close(fig)


def make_importance_plot(importance: pd.DataFrame, output_path: Path) -> None:
    selected = importance.head(7).sort_values("permutation_importance")
    fig, ax = plt.subplots(figsize=(7.2, 4.5))
    ax.barh(
        selected["feature"],
        selected["permutation_importance"],
        xerr=selected["std"],
        color="#a65383",
        alpha=0.85,
    )
    ax.set(
        xlabel="Decrease in held-out ROC-AUC after permutation",
        title="Permutation importance (selected model)",
    )
    ax.grid(axis="x", alpha=0.2)
    fig.tight_layout()
    fig.savefig(output_path, dpi=160)
    plt.close(fig)


def add_bullet(document: Document, text: str) -> None:
    document.add_paragraph(text, style="List Bullet")


def write_report(
    output_path: Path,
    metrics: dict[str, Any],
    comparison: pd.DataFrame,
    importance: pd.DataFrame,
    synthetic: bool,
    roc_plot: Path,
    importance_plot: Path,
) -> None:
    """Create a Word report containing the exact metrics from this model run."""
    document = Document()
    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10)

    document.add_heading("Predictive Modeling Strategy", 0)
    document.add_paragraph(
        "Beauty & wellness customer return prediction | Reproducible Python analysis"
    )
    document.add_heading("Executive summary", level=1)
    provenance = (
        "This run uses a seeded synthetic dataset, not company/customer records."
        if synthetic
        else "This run uses the supplied CSV; results depend on its data quality and representativeness."
    )
    document.add_paragraph(provenance)
    document.add_paragraph(
        f"The business target is whether a customer returns within 60 days of the "
        f"snapshot. The selected candidate is {metrics['selected_model']}, chosen "
        f"by mean five-fold training ROC-AUC. Its held-out ROC-AUC is "
        f"{metrics['test_metrics']['roc_auc']:.3f}, average precision is "
        f"{metrics['test_metrics']['pr_auc_average_precision']:.3f}, and "
        f"thresholded recall/precision are "
        f"{metrics['test_metrics']['recall']:.3f}/"
        f"{metrics['test_metrics']['precision']:.3f}."
    )

    document.add_heading("Business problem and target", level=1)
    document.add_paragraph(
        "A salon, spa, skincare clinic, or wellness studio can use a return-risk "
        "score to prioritize timely, relevant follow-up. The prediction unit is "
        "one customer at a defined snapshot. The target is returned_within_60d "
        "(1 if the customer made a subsequent visit within 60 days, otherwise 0). "
        "This is a supervised binary-classification task, not a causal estimate "
        "of whether a promotion will change behavior."
    )
    document.add_paragraph(
        "Features in this worked example are visits in the last 90 days, average "
        "ticket in USD, days since the last visit, membership status, share of "
        "visits using a discount, satisfaction score (1–5), and service category. "
        "Identifiers and post-snapshot data must not be model features."
    )

    document.add_heading("Data and reproducibility", level=1)
    document.add_paragraph(
        f"The dataset contains {metrics['n_rows']:,} rows; the positive return "
        f"rate is {metrics['positive_rate']:.1%}. The split is stratified "
        f"{metrics['train_rows']:,}/{metrics['test_rows']:,} "
        "(75% training / 25% held-out test, random seed 42). "
        "Model selection and threshold tuning use training data only; the test "
        "set is used once for final evaluation."
    )
    document.add_paragraph(
        "Synthetic-data limitation: generated relationships make this an "
        "executable worked example, not evidence of actual business performance. "
        "Replace the example with consented, properly time-aligned operational "
        "data before making customer-facing decisions."
    )

    document.add_heading("Candidate algorithms and rationale", level=1)
    document.add_heading("1. Logistic regression (interpretable baseline)", level=2)
    document.add_paragraph(
        "A regularized logistic classifier provides a transparent probability "
        "baseline. Numeric variables are standardized and categorical variables "
        "one-hot encoded in a leakage-safe pipeline. Class weighting addresses "
        "unequal class frequencies during fitting. Coefficient direction can "
        "describe associations, but it is not causal and can be unstable when "
        "features are correlated."
    )
    document.add_heading("2. Random forest (nonlinear benchmark)", level=2)
    document.add_paragraph(
        "A 300-tree random forest can capture nonlinear effects and feature "
        "interactions, such as recency having a different impact for members. "
        "The minimum leaf size of eight constrains complexity; balanced "
        "subsampling supports the minority class. It is less directly "
        "interpretable and may need probability calibration for operational use."
    )
    document.add_paragraph(
        "Both candidates use identical stratified folds, hold-out rows, and "
        "predictor columns. The candidate with the larger mean five-fold "
        "training ROC-AUC is selected automatically; a simple baseline is "
        "retained for comparison."
    )

    document.add_heading("Validation and measured results", level=1)
    document.add_paragraph(
        "Five-fold StratifiedKFold estimates training ROC-AUC and its standard "
        "deviation. Out-of-fold training probabilities select an F1-maximizing "
        "threshold from 0.10–0.80; the final test set is not used to tune it. "
        "The held-out report includes ROC-AUC, average precision (PR-AUC), "
        "precision, recall, F1, accuracy, specificity, and Brier score. "
        "ROC-AUC measures ranking; PR-AUC emphasizes positive-class retrieval; "
        "Brier score measures probability error (lower is better)."
    )
    table = document.add_table(rows=1, cols=4)
    table.style = "Light Shading Accent 1"
    for cell, value in zip(
        table.rows[0].cells,
        ["Candidate", "5-fold ROC-AUC (mean ± SD)", "Test ROC-AUC", "Test PR-AUC"],
    ):
        cell.text = value
    for _, row in comparison.iterrows():
        cells = table.add_row().cells
        cells[0].text = str(row["model"])
        cells[1].text = f"{row['cv_roc_auc_mean']:.3f} ± {row['cv_roc_auc_std']:.3f}"
        cells[2].text = f"{row['test_roc_auc']:.3f}"
        cells[3].text = f"{row['test_pr_auc']:.3f}"

    selected = metrics["test_metrics"]
    document.add_paragraph(
        f"Selected model: {metrics['selected_model']}. Training-only selected "
        f"threshold: {metrics['threshold']:.2f} (OOF F1={metrics['oof_f1']:.3f}). "
        f"On the held-out set, precision={selected['precision']:.3f}, "
        f"recall={selected['recall']:.3f}, F1={selected['f1']:.3f}, "
        f"specificity={selected['specificity']:.3f}, and "
        f"Brier={selected['brier_score']:.3f}. Confusion counts: "
        f"TN={selected['true_negative']}, FP={selected['false_positive']}, "
        f"FN={selected['false_negative']}, TP={selected['true_positive']}."
    )
    actual_returners = selected["true_positive"] + selected["false_negative"]
    actual_non_returners = selected["true_negative"] + selected["false_positive"]
    document.add_paragraph(
        f"Operational reading: at this return-class F1 threshold, the model "
        f"identified {selected['true_positive']} of {actual_returners} actual "
        f"returners, but correctly identified only {selected['true_negative']} "
        f"of {actual_non_returners} non-returners. It labels "
        f"{selected['true_positive'] + selected['false_positive']} of "
        f"{metrics['test_rows']} test customers as returners. Therefore this "
        "threshold is not a ready-made retention-campaign rule. For a "
        "re-engagement pilot, rank by lower predicted return probability and "
        "choose a contact cutoff using campaign capacity and economics; "
        "validate incremental impact with a randomized control group."
    )
    document.add_picture(str(roc_plot), width=Inches(6.2))
    document.paragraphs[-1].alignment = 1

    document.add_heading("Business interpretation and action", level=1)
    document.add_paragraph(
        "Use scores to rank customers for human-reviewed outreach, for example "
        "a service reminder or a relevant rebooking prompt. Start with a small "
        "pilot and compare incremental rebooking and margin against a randomized "
        "control group: predictive risk alone does not show that an offer works. "
        "Choose the operating threshold using the relative cost of missed "
        "rebookings, unnecessary contacts, and discount margin; the reported "
        "F1 threshold is only a neutral worked-example choice."
    )
    top_features = ", ".join(
        f"{row.feature} ({row.permutation_importance:.3f})"
        for row in importance.head(5).itertuples()
    )
    document.add_paragraph(
        "Permutation importance on held-out data ranks features by the drop in "
        f"ROC-AUC when shuffled. The leading example signals are {top_features}. "
        "They indicate predictive association, not causal drivers."
    )
    document.add_picture(str(importance_plot), width=Inches(6.2))
    document.paragraphs[-1].alignment = 1

    document.add_heading("Robustness, pitfalls, and mitigations", level=1)
    for item in [
        "Temporal leakage: construct every feature strictly before the prediction snapshot; use a later time period for a final deployment-style validation.",
        "Repeated customers: split by customer (and preferably time) when multiple snapshots per person exist so the same person cannot appear in train and test.",
        "Overfitting: compare training cross-validation with untouched test results, constrain tree leaf size, and tune only within training folds.",
        "Class imbalance and threshold drift: report PR-AUC and recall/precision, track prevalence, and re-select the threshold using business costs rather than defaulting to 0.50.",
        "Missingness and inconsistent logging: report data-quality checks, define missing-value treatment from training data, and monitor input distributions.",
        "Calibration and subgroup performance: check reliability/calibration and precision/recall by relevant service and membership groups before rollout; investigate gaps and avoid discriminatory targeting.",
        "Concept drift and intervention feedback: monitor performance over time; promotions can change the behavior being predicted, so evaluate campaigns with controlled experiments.",
    ]:
        add_bullet(document, item)

    document.add_heading("Implementation and next steps", level=1)
    document.add_paragraph(
        "The companion Python script writes the metrics JSON, row-level test "
        "predictions, permutation importances, plots, fitted model bundle, and "
        "this report. For real deployment, confirm the event window and feature "
        "availability, establish a time-based backtest, calibrate probabilities, "
        "review privacy/consent and fairness requirements, and monitor drift."
    )
    document.save(output_path)


def run(input_path: Path | None, output_dir: Path) -> dict[str, Any]:
    frame, synthetic = load_data(input_path)
    validate_data(frame)
    output_dir.mkdir(parents=True, exist_ok=True)
    if synthetic:
        frame.to_csv(output_dir / "synthetic_customer_returns.csv", index=False)

    x = frame[FEATURES]
    y = frame[TARGET].astype(int)
    models = build_models()
    from sklearn.inspection import permutation_importance
    from sklearn.metrics import average_precision_score, roc_auc_score
    from sklearn.model_selection import (
        StratifiedKFold,
        cross_val_predict,
        cross_val_score,
        train_test_split,
    )

    x_train, x_test, y_train, y_test = train_test_split(
        x,
        y,
        test_size=0.25,
        stratify=y,
        random_state=RANDOM_STATE,
    )
    folds = StratifiedKFold(n_splits=5, shuffle=True, random_state=RANDOM_STATE)
    comparison_rows: list[dict[str, Any]] = []
    fitted_models: dict[str, Any] = {}
    test_probabilities: dict[str, np.ndarray] = {}
    oof_probabilities: dict[str, np.ndarray] = {}

    for name, model in models.items():
        cv_scores = cross_val_score(
            model, x_train, y_train, cv=folds, scoring="roc_auc", n_jobs=1
        )
        oof = cross_val_predict(
            model,
            x_train,
            y_train,
            cv=folds,
            method="predict_proba",
            n_jobs=1,
        )[:, 1]
        model.fit(x_train, y_train)
        test_proba = model.predict_proba(x_test)[:, 1]
        oof_probabilities[name] = oof
        fitted_models[name] = model
        test_probabilities[name] = test_proba
        comparison_rows.append(
            {
                "model": name,
                "cv_roc_auc_mean": float(cv_scores.mean()),
                "cv_roc_auc_std": float(cv_scores.std(ddof=1)),
                "test_roc_auc": float(roc_auc_score(y_test, test_proba)),
                "test_pr_auc": float(average_precision_score(y_test, test_proba)),
            }
        )
        LOGGER.info(
            "%s: CV ROC-AUC %.3f ± %.3f; test ROC-AUC %.3f",
            name,
            cv_scores.mean(),
            cv_scores.std(ddof=1),
            roc_auc_score(y_test, test_proba),
        )

    comparison = pd.DataFrame(comparison_rows).sort_values(
        "cv_roc_auc_mean", ascending=False
    )
    selected_name = str(comparison.iloc[0]["model"])
    threshold, oof_f1 = choose_f1_threshold(
        y_train.to_numpy(), oof_probabilities[selected_name]
    )
    selected_model = fitted_models[selected_name]
    selected_probabilities = test_probabilities[selected_name]
    test_metrics = evaluate(y_test.to_numpy(), selected_probabilities, threshold)

    importance_result = permutation_importance(
        selected_model,
        x_test,
        y_test,
        scoring="roc_auc",
        n_repeats=10,
        random_state=RANDOM_STATE,
        n_jobs=-1,
    )
    importance = pd.DataFrame(
        {
            "feature": FEATURES,
            "permutation_importance": importance_result.importances_mean,
            "std": importance_result.importances_std,
        }
    ).sort_values("permutation_importance", ascending=False)

    roc_plot = output_dir / "roc_curves.png"
    importance_plot = output_dir / "feature_importance.png"
    make_roc_plot(y_test, test_probabilities, roc_plot)
    make_importance_plot(importance, importance_plot)
    comparison.to_csv(output_dir / "model_comparison.csv", index=False)
    importance.to_csv(output_dir / "feature_importance.csv", index=False)

    predictions = x_test.reset_index(drop=True).copy()
    predictions["actual_returned_within_60d"] = y_test.reset_index(drop=True)
    predictions["predicted_return_probability"] = selected_probabilities
    predictions["predicted_return_at_threshold"] = (
        selected_probabilities >= threshold
    ).astype(int)
    predictions.to_csv(output_dir / "test_predictions.csv", index=False)

    metrics: dict[str, Any] = {
        "synthetic_data": synthetic,
        "random_state": RANDOM_STATE,
        "n_rows": int(len(frame)),
        "train_rows": int(len(x_train)),
        "test_rows": int(len(x_test)),
        "positive_rate": float(y.mean()),
        "selected_model": selected_name,
        "threshold": threshold,
        "oof_f1": oof_f1,
        "test_metrics": test_metrics,
        "model_comparison": comparison.to_dict(orient="records"),
    }
    with (output_dir / "evaluation_metrics.json").open(
        "w", encoding="utf-8"
    ) as json_file:
        json.dump(metrics, json_file, indent=2)
    joblib.dump(
        {
            "model": selected_model,
            "threshold": threshold,
            "feature_columns": FEATURES,
            "target": TARGET,
            "synthetic_training_data": synthetic,
        },
        output_dir / "return_prediction_model.joblib",
    )
    write_report(
        output_dir / "predictive_modeling_strategy.docx",
        metrics,
        comparison,
        importance,
        synthetic,
        roc_plot,
        importance_plot,
    )
    LOGGER.info("Finished. Outputs saved to %s", output_dir.resolve())
    LOGGER.info(
        "Selected %s at threshold %.2f: ROC-AUC %.3f, precision %.3f, recall %.3f",
        selected_name,
        threshold,
        test_metrics["roc_auc"],
        test_metrics["precision"],
        test_metrics["recall"],
    )
    return metrics


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Train and evaluate beauty/wellness 60-day return models."
    )
    parser.add_argument(
        "--input",
        type=Path,
        help="Optional CSV with documented feature columns and returned_within_60d target.",
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=Path("outputs"),
        help="Directory for generated reports, metrics, plots, and model (default: outputs).",
    )
    parser.add_argument("--verbose", action="store_true", help="Enable progress logs.")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    logging.basicConfig(
        level=logging.INFO if args.verbose else logging.WARNING,
        format="%(levelname)s: %(message)s",
    )
    try:
        run(args.input, args.output_dir)
    except (OSError, RuntimeError, ValueError, TypeError, KeyError) as error:
        LOGGER.error("%s", error)
        raise SystemExit(1) from error


if __name__ == "__main__":
    main()
